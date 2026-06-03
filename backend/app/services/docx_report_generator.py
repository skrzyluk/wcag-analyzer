"""
docx_report_generator.py - Generator raportu audytu WCAG 2.1 AA w formacie Word (.docx).
"""

from __future__ import annotations

import io
import datetime
from typing import TYPE_CHECKING

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ── Kolory ─────────────────────────────────────────────────────────────────

_COLOR_PRIMARY   = RGBColor(0x1F, 0x4E, 0x78)   # ciemny niebieski
_COLOR_CRITICAL  = RGBColor(0xC0, 0x00, 0x00)   # czerwony
_COLOR_MAJOR     = RGBColor(0xC5, 0x5A, 0x11)   # pomarańczowy
_COLOR_MINOR     = RGBColor(0x37, 0x56, 0x23)   # zielony
_COLOR_WARNING   = RGBColor(0x7F, 0x60, 0x00)   # żółto-brązowy
_COLOR_WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
_COLOR_LIGHT_BG  = RGBColor(0xD6, 0xDC, 0xE4)
_COLOR_HEADER_BG = "1F4E78"   # hex string dla fill komórek
_COLOR_CRIT_BG   = "FCE4D6"
_COLOR_MAJOR_BG  = "FFF2CC"
_COLOR_MINOR_BG  = "E2EFDA"


# ── Helpers ─────────────────────────────────────────────────────────────────

def _set_cell_bg(cell, hex_color: str) -> None:
    """Ustaw kolor tła komórki tabeli (przez XML)."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _bold_run(para, text: str, color: RGBColor | None = None, size_pt: int = 11) -> None:
    run = para.add_run(text)
    run.bold = True
    run.font.size = Pt(size_pt)
    if color:
        run.font.color.rgb = color


def _normal_run(para, text: str, size_pt: int = 10) -> None:
    run = para.add_run(text)
    run.font.size = Pt(size_pt)


def _severity_label(severity: str) -> str:
    return {
        "critical": "Krytyczne",
        "major":    "Poważne",
        "minor":    "Drobne",
        "warning":  "Ostrzeżenia",
    }.get(severity, severity.capitalize())


def _severity_color(severity: str) -> RGBColor:
    return {
        "critical": _COLOR_CRITICAL,
        "major":    _COLOR_MAJOR,
        "minor":    _COLOR_MINOR,
        "warning":  _COLOR_WARNING,
    }.get(severity, _COLOR_PRIMARY)


def _severity_bg(severity: str) -> str:
    return {
        "critical": _COLOR_CRIT_BG,
        "major":    _COLOR_MAJOR_BG,
        "minor":    _COLOR_MINOR_BG,
        "warning":  "FFF9C4",
    }.get(severity, "FFFFFF")


# ── Główna klasa ─────────────────────────────────────────────────────────────

class DOCXReportGenerator:
    """
    Generuje raport audytu WCAG 2.1 AA jako plik Word (.docx).

    Użycie:
        gen = DOCXReportGenerator()
        docx_bytes = gen.generate_bytes(report_dict)
        gen.generate_file(report_dict, "raport.docx")
    """

    def generate_bytes(self, report: dict) -> bytes:
        doc = self._build(report)
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf.read()

    def generate_file(self, report: dict, path: str) -> str:
        doc = self._build(report)
        doc.save(path)
        return path

    # ── Budowanie dokumentu ─────────────────────────────────────────────────

    def _build(self, report: dict) -> Document:
        doc = Document()
        self._set_margins(doc)
        self._add_title(doc, report)
        self._add_metadata(doc, report)
        self._add_summary(doc, report)
        self._add_issues(doc, report)
        self._add_footer_note(doc)
        return doc

    def _set_margins(self, doc: Document) -> None:
        for section in doc.sections:
            section.top_margin    = Cm(2)
            section.bottom_margin = Cm(2)
            section.left_margin   = Cm(2.5)
            section.right_margin  = Cm(2)

    # ── Tytuł ───────────────────────────────────────────────────────────────

    def _add_title(self, doc: Document, report: dict) -> None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("Raport audytu dostępności WCAG 2.1 AA")
        run.bold = True
        run.font.size = Pt(18)
        run.font.color.rgb = _COLOR_PRIMARY

        doc.add_paragraph()  # odstęp

    # ── Metadane ────────────────────────────────────────────────────────────

    def _add_metadata(self, doc: Document, report: dict) -> None:
        filename = report.get("filename", "—")
        date_raw = report.get("analysis_date", "")
        try:
            dt = datetime.datetime.fromisoformat(date_raw.rstrip("Z"))
            date_str = dt.strftime("%d.%m.%Y %H:%M")
        except Exception:
            date_str = date_raw or "—"

        page_url = report.get("page_url", "")
        page_title = report.get("page_title", "")

        table = doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.LEFT

        def _meta_row(label: str, value: str) -> None:
            row = table.add_row()
            row.cells[0].text = label
            row.cells[0].paragraphs[0].runs[0].bold = True
            row.cells[0].paragraphs[0].runs[0].font.size = Pt(10)
            row.cells[1].text = value
            row.cells[1].paragraphs[0].runs[0].font.size = Pt(10)
            _set_cell_bg(row.cells[0], "D6DCE4")

        _meta_row("Plik źródłowy", filename)
        _meta_row("Data analizy", date_str)
        if page_title:
            _meta_row("Tytuł strony", page_title)
        if page_url:
            _meta_row("URL strony", page_url)

        table.columns[0].width = Cm(5)
        table.columns[1].width = Cm(13)

        doc.add_paragraph()

    # ── Podsumowanie ────────────────────────────────────────────────────────

    def _add_summary(self, doc: Document, report: dict) -> None:
        h = doc.add_heading("Podsumowanie wyników", level=1)
        h.runs[0].font.color.rgb = _COLOR_PRIMARY

        summary = report.get("summary", {})
        compliance = summary.get("compliance_percentage", 0)
        total      = summary.get("total_issues", 0)
        critical   = summary.get("critical_errors", 0)
        major      = summary.get("major_errors", 0)
        minor      = summary.get("minor_errors", 0)
        warnings   = summary.get("warnings", 0)
        passed     = summary.get("passed_checks", 0)
        checks     = summary.get("total_checks", 50)

        # Wynik zgodności
        p = doc.add_paragraph()
        p.add_run("Zgodność z WCAG 2.1 AA: ").bold = True
        run = p.add_run(f"{compliance}%")
        run.bold = True
        run.font.size = Pt(14)
        run.font.color.rgb = _COLOR_CRITICAL if compliance < 50 else (
            _COLOR_MAJOR if compliance < 80 else _COLOR_MINOR
        )

        # Tabela statystyk
        table = doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"

        def _stat_row(label: str, value, bg: str = "FFFFFF", bold_val: bool = False) -> None:
            row = table.add_row()
            row.cells[0].text = label
            row.cells[0].paragraphs[0].runs[0].font.size = Pt(10)
            _set_cell_bg(row.cells[0], "D6DCE4")
            row.cells[1].text = str(value)
            row.cells[1].paragraphs[0].runs[0].font.size = Pt(10)
            if bold_val:
                row.cells[1].paragraphs[0].runs[0].bold = True
            _set_cell_bg(row.cells[1], bg)

        _stat_row("Sprawdzonych kryteriów", checks)
        _stat_row("Kryteria spełnione", passed, bg="E2EFDA", bold_val=True)
        _stat_row("Łączna liczba problemów", total, bold_val=True)
        _stat_row("Krytyczne", critical, bg=_COLOR_CRIT_BG, bold_val=critical > 0)
        _stat_row("Poważne", major, bg=_COLOR_MAJOR_BG, bold_val=major > 0)
        _stat_row("Drobne", minor, bg=_COLOR_MINOR_BG)
        if warnings:
            _stat_row("Ostrzeżenia", warnings)

        table.columns[0].width = Cm(8)
        table.columns[1].width = Cm(4)

        doc.add_paragraph()

    # ── Lista problemów ─────────────────────────────────────────────────────

    def _add_issues(self, doc: Document, report: dict) -> None:
        issues = report.get("issues", [])
        if not issues:
            h = doc.add_heading("Wyniki analizy", level=1)
            h.runs[0].font.color.rgb = _COLOR_PRIMARY
            doc.add_paragraph("Nie wykryto żadnych problemów dostępności.")
            return

        h = doc.add_heading("Wyniki analizy", level=1)
        h.runs[0].font.color.rgb = _COLOR_PRIMARY

        # Grupuj po severity
        order = ["critical", "major", "minor", "warning"]
        groups: dict[str, list] = {s: [] for s in order}
        for issue in issues:
            sev = issue.get("severity", "minor") if isinstance(issue, dict) else getattr(issue, "severity", "minor")
            if sev in groups:
                groups[sev].append(issue)
            else:
                groups["minor"].append(issue)

        for severity in order:
            group_issues = groups[severity]
            if not group_issues:
                continue

            label = _severity_label(severity)
            color = _severity_color(severity)
            bg    = _severity_bg(severity)

            h2 = doc.add_heading(f"{label} ({len(group_issues)})", level=2)
            h2.runs[0].font.color.rgb = color

            for idx, issue in enumerate(group_issues, 1):
                if isinstance(issue, dict):
                    criterion    = issue.get("wcag_criterion", "—")
                    description  = issue.get("description", "—")
                    recommendation = issue.get("recommendation", "—")
                    affected     = issue.get("affected_elements", [])
                else:
                    criterion    = getattr(issue, "wcag_criterion", "—")
                    description  = getattr(issue, "description", "—")
                    recommendation = getattr(issue, "recommendation", "—")
                    affected     = getattr(issue, "affected_elements", [])

                # Tabela dla jednego problemu
                table = doc.add_table(rows=0, cols=2)
                table.style = "Table Grid"

                # Nagłówek wiersza z numerem i kryterium
                row = table.add_row()
                row.cells[0].merge(row.cells[1])
                row.cells[0].text = f"{idx}. {criterion}"
                r = row.cells[0].paragraphs[0].runs[0]
                r.bold = True
                r.font.size = Pt(10)
                r.font.color.rgb = color
                _set_cell_bg(row.cells[0], bg)

                def _issue_row(label: str, text: str) -> None:
                    r2 = table.add_row()
                    r2.cells[0].text = label
                    r2.cells[0].paragraphs[0].runs[0].bold = True
                    r2.cells[0].paragraphs[0].runs[0].font.size = Pt(9)
                    _set_cell_bg(r2.cells[0], "D6DCE4")
                    r2.cells[1].text = text or "—"
                    r2.cells[1].paragraphs[0].runs[0].font.size = Pt(9)

                _issue_row("Problem", description)
                _issue_row("Rekomendacja", recommendation)

                # Fragmenty kodu (pierwsze 3 elementy)
                snippets = []
                for el in (affected or [])[:3]:
                    if isinstance(el, dict):
                        s = el.get("html_snippet", "")
                    else:
                        s = getattr(el, "html_snippet", "")
                    if s:
                        snippets.append(s)

                if snippets:
                    r2 = table.add_row()
                    r2.cells[0].text = "Fragment HTML"
                    r2.cells[0].paragraphs[0].runs[0].bold = True
                    r2.cells[0].paragraphs[0].runs[0].font.size = Pt(9)
                    _set_cell_bg(r2.cells[0], "D6DCE4")
                    snip_run = r2.cells[1].paragraphs[0].add_run("\n".join(snippets))
                    snip_run.font.name = "Consolas"
                    snip_run.font.size = Pt(8)

                table.columns[0].width = Cm(4)
                table.columns[1].width = Cm(13)

                doc.add_paragraph()  # odstęp między problemami

    # ── Stopka ──────────────────────────────────────────────────────────────

    def _add_footer_note(self, doc: Document) -> None:
        doc.add_paragraph()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(
            f"Raport wygenerowany automatycznie przez WCAG 2.1 AA Analyzer "
            f"· {datetime.datetime.now().strftime('%d.%m.%Y %H:%M')}"
        )
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        run.italic = True
