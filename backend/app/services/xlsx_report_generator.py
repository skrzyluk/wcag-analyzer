"""
xlsx_report_generator.py - Budowanie pliku .xlsx z wynikami audytu WCAG 2.1 AA.

Implementacja zgodna ze specyfikacja INSTRUKCJA_RAPORT_XLSX.md (wzor 1:1
z pliku Audyt_WCAG_2_1_AA_PRZYKLAD.xlsx).

Struktura skoroszytu:
  1..N  arkusze "tabela naruszen" (jeden na audytowana podstrone)
  N+1   arkusz "Kolejnosc fokusa" (opcjonalny)
  N+2   arkusz "Legenda" (statyczny boilerplate)
"""

from __future__ import annotations

import io
from typing import TYPE_CHECKING

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.services.issue_mapper import ExcelRow, PageResult, SeparatorRow
from app.services.focus_order_analyzer import FocusStep

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet

# ---------------------------------------------------------------------------
# Paleta kolorow (RRGGBB) — wg INSTRUKCJA_RAPORT_XLSX.md §2
# ---------------------------------------------------------------------------

C_HEADER_BG    = "1F4E78"   # granat — tytul / naglowki kolumn
C_HEADER_FG    = "FFFFFF"   # bialy tekst na granacie
C_DESC_BG      = "F2F2F2"   # tlo wiersza opisu
C_SEPARATOR_BG = "595959"   # szary — separator sekcji
C_SEPARATOR_FG = "FFFFFF"
C_ZEBRA        = "F5F5F5"   # pasek zebry (co drugi wiersz)
C_WHITE        = "FFFFFF"
C_BORDER       = "CCCCCC"   # cienkie obramowanie

STATUS_COLORS = {
    "KRYTYCZNE": "F4B084",
    "NOWE":      "E2EFDA",
    "WSPOLNE":   "FFF2CC",
    "OK":        "C6EFCE",
}

LEVEL_COLORS = {
    "A":   "FCE4D6",
    "AA":  "FFF2CC",
    "AAA": "DDEBF7",
}

# ---------------------------------------------------------------------------
# Typografia
# ---------------------------------------------------------------------------

FONT_DEFAULT = "Arial"
FONT_CODE    = "Consolas"

# Indeksy kolumn (1-based) z czcionka kodu: J = Fragment kodu, K = do zmiany
CODE_COLS = {10, 11}

# ---------------------------------------------------------------------------
# Tabela naruszen — kolumny i szerokosci (§4.1)
# ---------------------------------------------------------------------------

COLUMN_HEADERS = [
    "ID", "Status", "Obszar", "Element", "Kryterium", "Poziom",
    "Opis ze standardu", "Problem", "Naprawa",
    "Fragment kodu", "Fragment kodu do zmiany",
]
# Szerokosci wg wzoru. Kolumny I (9) i K (11) celowo zostawione domyslne.
COLUMN_WIDTHS = {1: 10, 2: 11, 3: 16, 4: 30, 5: 22, 6: 7, 7: 42, 8: 46, 10: 36}

# Mapowanie ID sekcji -> etykieta poziomu (do tekstu separatora)
SECTION_LEVEL = {"R1": "A", "R2": "AA"}

# ---------------------------------------------------------------------------
# Arkusz "Kolejnosc fokusa" — kolumny i szerokosci (§5.1)
# ---------------------------------------------------------------------------

FOCUS_COLUMN_HEADERS = ["Krok", "Obszar strony", "Element / kontrolka", "Uzasadnienie"]
FOCUS_COLUMN_WIDTHS  = {1: 8, 2: 20, 3: 40, 4: 60}

# ---------------------------------------------------------------------------
# Wysokosci wierszy (wg wzoru)
# ---------------------------------------------------------------------------

H_TITLE     = 27.8
H_DESC      = 36.0
H_HEADER    = 19.5
H_SEPARATOR = 18.0
H_DATA      = 51.8

# ---------------------------------------------------------------------------
# Legenda — statyczny boilerplate (§6)
# ---------------------------------------------------------------------------

LEGEND_TITLE = "Legenda raportu audytowego — Audyt WCAG 2.1 AA"

# Kazdy blok: naglowek (3 komorki) + wiersze danych.
#   kind="color" -> caly wiersz wypelniony kolorem, B = kod hex
#   kind="plain" -> A,B tlo szare (F5F5F5), C bez tla
LEGEND_BLOCKS = [
    {
        "kind": "color",
        "header": ("Status wiersza", "Kolor", "Znaczenie"),
        "rows": [
            ("KRYTYCZNE", STATUS_COLORS["KRYTYCZNE"],
             "Problem o bardzo wysokim wplywie na uzytkownikow. Wymaga "
             "natychmiastowej naprawy — czesto blokuje funkcjonalnosc dla "
             "uzytkownikow czytnikow ekranu lub klawiatury."),
            ("NOWE", STATUS_COLORS["NOWE"],
             "Problem wykryty przy analizie konkretnej podstrony lub komponentu. "
             "Specyficzny dla danego obszaru aplikacji."),
            ("WSPOLNE", STATUS_COLORS["WSPOLNE"],
             "Problem wystepujacy na wielu podstronach. Przy naprawie wystarczy "
             "zmienic jedno miejsce (np. szablon), aby rozwiazac go wszedzie."),
            ("OK", STATUS_COLORS["OK"],
             "Element zaimplementowany zgodnie ze standardem — warty "
             "udokumentowania jako wzorzec do nasladowania."),
        ],
    },
    {
        "kind": "color",
        "header": ("Poziom WCAG", "Kolor", "Znaczenie"),
        "rows": [
            ("A", LEVEL_COLORS["A"],
             "Najnizszy poziom zgodnosci. Kryteria absolutnie konieczne — ich "
             "naruszenie czyni tresc niedostepna dla duzych grup uzytkownikow."),
            ("AA", LEVEL_COLORS["AA"],
             "Sredni poziom, standard dla sektora publicznego. Polska ustawa o "
             "dostepnosci cyfrowej wymaga zgodnosci z WCAG 2.1 AA."),
            ("AAA", LEVEL_COLORS["AAA"],
             "Najwyzszy poziom, czesto niemozliwy do osiagniecia w calosci. "
             "Oznaczamy wylacznie gdy kryterium jest istotne kontekstowo."),
        ],
    },
    {
        "kind": "plain",
        "header": ("Kolumna raportu", "Dla kogo", "Zawartosc"),
        "rows": [
            ("ID", "Audytor / PM",
             "Unikalny identyfikator wiersza problemu. Format: "
             "{PREFIKS}-R{runda}-{nr}, np. KPLT-R1-001. R1 = naruszenia "
             "poziomu A, R2 = naruszenia poziomu AA."),
            ("Status", "Audytor",
             "Typ problemu: KRYTYCZNE / NOWE / WSPOLNE / OK. Pozwala szybko "
             "filtrowac priorytety."),
            ("Obszar", "Audytor / PM",
             "Jedna z 4 zasad WCAG: Postrzegalnosc / Funkcjonalnosc / "
             "Zrozumialosc / Solidnosc."),
            ("Element", "Analityk / Tester",
             "Konkretny element interfejsu, ktorego dotyczy problem — z "
             "informacja o podstronie."),
            ("Kryterium", "Audytor / Ekspert WCAG",
             "Numer i nazwa kryterium sukcesu wg WCAG 2.1. Uzywane do mapowania "
             "na wymagania prawne."),
            ("Poziom", "Audytor / PM",
             "A / AA / AAA. Poziom zgodnosci wedlug WCAG 2.1."),
            ("Opis ze standardu", "Analityk / PM",
             "Oficjalne brzmienie kryterium sukcesu (tlumaczenie robocze). "
             "Pozwala zrozumiec intencje standardu."),
            ("Problem", "Analityk / Tester",
             "Wyjasnienie problemu: co sie dzieje, dlaczego to blad, kogo "
             "dotyczy, jakie sa skutki."),
            ("Naprawa", "Programista",
             "Konkretne instrukcje naprawy — krok po kroku, z odniesieniem do "
             "atrybutow HTML/ARIA, klas CSS lub wzorcow WAI-ARIA."),
            ("Fragment kodu", "Programista",
             "Aktualna, problematyczna implementacja (czcionka Consolas)."),
            ("Fragment kodu do zmiany", "Programista",
             "Propozycja poprawionej implementacji (czcionka Consolas)."),
        ],
    },
    {
        "kind": "plain",
        "header": ("Inne skroty", "Dla kogo", "Zawartosc"),
        "rows": [
            ("Uzytkownik AT", "Wszyscy",
             "AT = Assistive Technology (technologia wspomagajaca). Uzytkownik AT "
             "to osoba korzystajaca z narzedzi dostepnosci: czytnik ekranu "
             "(NVDA, JAWS, VoiceOver), klawiatura jako jedyne urzadzenie wejscia, "
             "wyswietlacz brajlowski, powiekszanie ekranu, sterowanie glosem."),
        ],
    },
]

LEGEND_COL_WIDTHS = {1: 26, 3: 90}  # B celowo domyslne (wg wzoru)


# ---------------------------------------------------------------------------
# Pomocnicze fabryki stylow
# ---------------------------------------------------------------------------

def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _font(bold=False, italic=False, color="000000", name=FONT_DEFAULT, size=9) -> Font:
    return Font(bold=bold, italic=italic, color=color, name=name, size=size)


def _align(wrap=True, horizontal="left", vertical="top") -> Alignment:
    return Alignment(wrap_text=wrap, horizontal=horizontal, vertical=vertical)


def _border() -> Border:
    s = Side(style="thin", color=C_BORDER)
    return Border(left=s, right=s, top=s, bottom=s)


# ---------------------------------------------------------------------------
# Generator
# ---------------------------------------------------------------------------

class XLSXReportGenerator:
    """
    Buduje plik .xlsx z wynikami audytu WCAG 2.1 AA.

        generate_bytes(pages, focus_steps=None) -> bytes
        generate_file(pages, path, focus_steps=None) -> str
    """

    def __init__(self):
        self.wb = openpyxl.Workbook()
        self.wb.remove(self.wb.active)

    # -- API publiczne --------------------------------------------------

    def generate_bytes(
        self,
        pages: list[PageResult],
        focus_steps: list[FocusStep] | None = None,
    ) -> bytes:
        self._build(pages, focus_steps)
        buf = io.BytesIO()
        self.wb.save(buf)
        buf.seek(0)
        return buf.read()

    def generate_file(
        self,
        pages: list[PageResult],
        path: str,
        focus_steps: list[FocusStep] | None = None,
    ) -> str:
        self._build(pages, focus_steps)
        self.wb.save(path)
        return path

    # -- Budowa skoroszytu ----------------------------------------------

    def _build(
        self,
        pages: list[PageResult],
        focus_steps: list[FocusStep] | None,
    ) -> None:
        for page in pages:
            self._create_analysis_sheet(page)
        if focus_steps is not None:
            self._create_focus_order_sheet(focus_steps)
        self._create_legend_sheet()

    # ------------------------------------------------------------------
    # Arkusz "tabela naruszen"
    # ------------------------------------------------------------------

    def _create_analysis_sheet(self, page: PageResult) -> None:
        ws = self.wb.create_sheet(title=page.sheet_name[:31])

        # Wiersz 1 — tytul (scalony A1:K1)
        ws.merge_cells("A1:K1")
        c = ws["A1"]
        c.value = page.title
        c.font  = _font(bold=True, color=C_HEADER_FG, size=12)
        c.fill  = _fill(C_HEADER_BG)
        c.alignment = _align(wrap=True, horizontal="left", vertical="center")
        ws.row_dimensions[1].height = H_TITLE

        # Wiersz 2 — opis (scalony A2:K2)
        ws.merge_cells("A2:K2")
        c = ws["A2"]
        c.value = page.description
        c.font  = _font(bold=True)
        c.fill  = _fill(C_DESC_BG)
        c.alignment = _align(wrap=True, horizontal="left", vertical="center")
        ws.row_dimensions[2].height = H_DESC

        # Wiersz 3 — naglowki kolumn
        for col_idx, header in enumerate(COLUMN_HEADERS, 1):
            c = ws.cell(row=3, column=col_idx, value=header)
            c.font      = _font(bold=True, color=C_HEADER_FG)
            c.fill      = _fill(C_HEADER_BG)
            c.alignment = _align(wrap=True, horizontal="center", vertical="center")
            c.border    = _border()
        ws.row_dimensions[3].height = H_HEADER

        ws.freeze_panes = "A4"
        ws.auto_filter.ref = f"A3:{get_column_letter(len(COLUMN_HEADERS))}3"

        # Wiersze danych + separatory (zebra resetowana per sekcja)
        row_num = 4
        zebra_idx = 0
        for item in page.rows:
            if isinstance(item, SeparatorRow):
                row_num = self._write_separator(ws, row_num, item)
                zebra_idx = 0  # reset zebry na poczatku sekcji
            elif isinstance(item, ExcelRow):
                row_num = self._write_data_row(ws, row_num, item, zebra_idx)
                zebra_idx += 1

        self._set_column_widths(ws, COLUMN_WIDTHS)

    def _write_separator(self, ws: "Worksheet", row_num: int, sep: SeparatorRow) -> int:
        level = SECTION_LEVEL.get(sep.section_id, sep.section_name.upper())
        ws.merge_cells(f"A{row_num}:K{row_num}")
        c = ws.cell(row=row_num, column=1, value=f"── NARUSZENIA POZIOMU {level} ──")
        c.font      = _font(bold=True, italic=True, color=C_SEPARATOR_FG)
        c.fill      = _fill(C_SEPARATOR_BG)
        c.alignment = _align(wrap=False, horizontal="left", vertical="center")
        ws.row_dimensions[row_num].height = H_SEPARATOR
        return row_num + 1

    def _write_data_row(
        self, ws: "Worksheet", row_num: int, row: ExcelRow, zebra_idx: int
    ) -> int:
        zebra_bg = C_ZEBRA if zebra_idx % 2 == 1 else C_WHITE
        status_bg = STATUS_COLORS.get(row.status, zebra_bg)
        level_bg  = LEVEL_COLORS.get(row.poziom, zebra_bg)

        for col_idx, value in enumerate(row.as_tuple(), 1):
            c = ws.cell(row=row_num, column=col_idx, value=value)
            font_name = FONT_CODE if col_idx in CODE_COLS else FONT_DEFAULT
            c.font      = _font(name=font_name, bold=(col_idx == 2))
            c.alignment = _align(wrap=True, horizontal="left", vertical="top")
            c.border    = _border()
            if col_idx == 2:       # Status
                c.fill = _fill(status_bg)
            elif col_idx == 6:     # Poziom
                c.fill = _fill(level_bg)
            else:
                c.fill = _fill(zebra_bg)

        ws.row_dimensions[row_num].height = H_DATA
        return row_num + 1

    @staticmethod
    def _set_column_widths(ws: "Worksheet", widths: dict[int, float]) -> None:
        for col_idx, width in widths.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = width

    # ------------------------------------------------------------------
    # Arkusz "Kolejnosc fokusa"
    # ------------------------------------------------------------------

    def _create_focus_order_sheet(self, steps: list[FocusStep]) -> None:
        ws = self.wb.create_sheet(title="Kolejnosc fokusa")

        ws.merge_cells("A1:D1")
        c = ws["A1"]
        c.value = "Rekomendowana kolejnosc przelaczania fokusa (klawisz Tab)"
        c.font      = _font(bold=True, color=C_HEADER_FG, size=12)
        c.fill      = _fill(C_HEADER_BG)
        c.alignment = _align(wrap=True, horizontal="left", vertical="center")
        ws.row_dimensions[1].height = H_TITLE

        ws.merge_cells("A2:D2")
        c = ws["A2"]
        c.value = (
            "Rekomendacja oparta na strukturze DOM i semantyce HTML. Kolejnosc "
            "powinna byc logiczna i zgodna z ukladem wizualnym — WCAG 2.4.3 (A) "
            "oraz 2.4.1 (Mozliwosc pominiecia blokow)."
        )
        c.font      = _font(bold=True)
        c.fill      = _fill(C_DESC_BG)
        c.alignment = _align(wrap=True, horizontal="left", vertical="center")
        ws.row_dimensions[2].height = 30

        for col_idx, header in enumerate(FOCUS_COLUMN_HEADERS, 1):
            c = ws.cell(row=3, column=col_idx, value=header)
            c.font      = _font(bold=True, color=C_HEADER_FG)
            c.fill      = _fill(C_HEADER_BG)
            c.alignment = _align(wrap=True, horizontal="center", vertical="center")
            c.border    = _border()
        ws.row_dimensions[3].height = H_HEADER

        ws.freeze_panes = "A4"
        ws.auto_filter.ref = "A3:D3"

        row_num = 4
        if steps:
            # Separator sekcji
            ws.merge_cells(f"A{row_num}:D{row_num}")
            c = ws.cell(row=row_num, column=1, value="── KOLEJNOSC FOKUSA ──")
            c.font      = _font(bold=True, italic=True, color=C_SEPARATOR_FG)
            c.fill      = _fill(C_SEPARATOR_BG)
            c.alignment = _align(wrap=False, horizontal="left", vertical="center")
            ws.row_dimensions[row_num].height = H_SEPARATOR
            row_num += 1

            for idx, step in enumerate(steps):
                zebra_bg = C_ZEBRA if idx % 2 == 1 else C_WHITE
                values = [step.krok, step.obszar_strony, step.element, step.uzasadnienie]
                for col_idx, val in enumerate(values, 1):
                    c = ws.cell(row=row_num, column=col_idx, value=val)
                    c.font   = _font(bold=(col_idx == 1))
                    c.fill   = _fill(zebra_bg)
                    if col_idx == 1:
                        c.alignment = _align(wrap=True, horizontal="center", vertical="top")
                    else:
                        c.alignment = _align(wrap=True, horizontal="left", vertical="top")
                    c.border = _border()
                ws.row_dimensions[row_num].height = 36
                row_num += 1
        else:
            ws.merge_cells(f"A{row_num}:D{row_num}")
            c = ws.cell(row=row_num, column=1,
                        value="Brak danych — analiza fokusa wymaga sparsowanego HTML.")
            c.font      = _font(italic=True, color="888888")
            c.alignment = _align(wrap=False, horizontal="left", vertical="center")
            ws.row_dimensions[row_num].height = 20

        self._set_column_widths(ws, FOCUS_COLUMN_WIDTHS)

    # ------------------------------------------------------------------
    # Arkusz "Legenda" (statyczny boilerplate)
    # ------------------------------------------------------------------

    def _create_legend_sheet(self) -> None:
        ws = self.wb.create_sheet(title="Legenda")

        # Wiersz 1 — tytul (scalony A1:C1)
        ws.merge_cells("A1:C1")
        c = ws["A1"]
        c.value = LEGEND_TITLE
        c.font  = _font(bold=True, color=C_HEADER_FG, size=12)
        c.fill  = _fill(C_HEADER_BG)
        c.alignment = _align(wrap=True, horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 25.5

        row_num = 2
        for block in LEGEND_BLOCKS:
            row_num = self._write_legend_header(ws, row_num, block["header"])
            for entry in block["rows"]:
                row_num = self._write_legend_row(ws, row_num, block["kind"], entry)

        self._set_column_widths(ws, LEGEND_COL_WIDTHS)

    def _write_legend_header(self, ws: "Worksheet", row_num: int, header: tuple) -> int:
        for col_idx, text in enumerate(header, 1):
            c = ws.cell(row=row_num, column=col_idx, value=text)
            c.font      = _font(bold=True, color=C_HEADER_FG)
            c.fill      = _fill(C_HEADER_BG)
            c.alignment = _align(wrap=False, horizontal="center", vertical="center")
            c.border    = _border()
        ws.row_dimensions[row_num].height = H_HEADER
        return row_num + 1

    def _write_legend_row(
        self, ws: "Worksheet", row_num: int, kind: str, entry: tuple
    ) -> int:
        label, mid, desc = entry
        for col_idx, value in enumerate((label, mid, desc), 1):
            c = ws.cell(row=row_num, column=col_idx, value=value)
            if kind == "color":
                c.fill = _fill(mid)  # caly wiersz w kolorze statusu/poziomu
            elif col_idx <= 2:
                c.fill = _fill(C_ZEBRA)
            c.font = _font(bold=(col_idx == 1))
            horizontal = "center" if (kind == "color" and col_idx <= 2) else "left"
            c.alignment = _align(wrap=True, horizontal=horizontal, vertical="top")
            c.border = _border()
        ws.row_dimensions[row_num].height = 48
        return row_num + 1
