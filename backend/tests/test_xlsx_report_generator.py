"""
test_xlsx_report_generator.py - Testy jednostkowe dla XLSXReportGenerator.

Uruchamianie:
    cd backend
    pytest tests/test_xlsx_report_generator.py -v
"""

import io
import pytest
import openpyxl
from openpyxl.styles import PatternFill

from app.services.issue_mapper import IssueMapper, Issue, AffectedElement, PageResult
from app.services.focus_order_analyzer import FocusStep
from app.services.xlsx_report_generator import (
    XLSXReportGenerator,
    STATUS_COLORS,
    LEVEL_COLORS,
    COLUMN_HEADERS,
    FOCUS_COLUMN_HEADERS,
    C_HEADER_BG,
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def make_workbook(pages, focus_steps=None):
    gen = XLSXReportGenerator()
    data = gen.generate_bytes(pages, focus_steps=focus_steps)
    return openpyxl.load_workbook(io.BytesIO(data))


def make_page(issues=None, sheet_name="Test", prefix="TST1", title="Test"):
    if issues is None:
        issues = [
            Issue(id="i1", wcag_criterion="1.1.1 Tresc", severity="critical",
                  description="Brak alt.", recommendation="Dodaj alt.",
                  affected_elements=[AffectedElement(tag="img",
                      html_snippet="<img>", html_after='<img alt="ok">')]),
            Issue(id="i2", wcag_criterion="1.4.3 Kontrast", severity="major",
                  description="Niski kontrast.", recommendation="Zwieksz kontrast.",
                  affected_elements=[]),
        ]
    mapper = IssueMapper()
    return mapper.map_page(issues=issues, sheet_name=sheet_name,
                           prefix=prefix, title=title)


def make_focus_steps():
    return [
        FocusStep(krok="Pre", obszar_strony="Naglowek strony",
                  element="Link: Skip", uzasadnienie="WCAG 2.4.1"),
        FocusStep(krok="1", obszar_strony="Nawigacja",
                  element="Link: Home", uzasadnienie="WCAG 2.4.4"),
        FocusStep(krok="2", obszar_strony="Glowna tresc",
                  element="Przycisk: Wyslij", uzasadnienie="WCAG 4.1.2"),
    ]


def cell_fill_hex(cell) -> str:
    fill = cell.fill
    if fill and fill.fgColor and fill.fgColor.type == "rgb":
        return fill.fgColor.rgb[-6:]  # pomijamy alpha (AARRGGBB)
    return ""


# ---------------------------------------------------------------------------
# Testy: struktura workbooka
# ---------------------------------------------------------------------------

class TestWorkbookStructure:
    def test_single_analysis_sheet_created(self):
        wb = make_workbook([make_page(sheet_name="Logowanie")])
        assert "Logowanie" in wb.sheetnames

    def test_legend_sheet_always_present(self):
        wb = make_workbook([make_page()])
        assert "Legenda" in wb.sheetnames

    def test_focus_sheet_present_when_steps_given(self):
        wb = make_workbook([make_page()], focus_steps=make_focus_steps())
        assert "Kolejnosc fokusa" in wb.sheetnames

    def test_focus_sheet_absent_when_none(self):
        wb = make_workbook([make_page()], focus_steps=None)
        assert "Kolejnosc fokusa" not in wb.sheetnames

    def test_focus_sheet_present_when_empty_list(self):
        wb = make_workbook([make_page()], focus_steps=[])
        assert "Kolejnosc fokusa" in wb.sheetnames

    def test_sheet_order_analysis_focus_legend(self):
        wb = make_workbook([make_page(sheet_name="Analiza")],
                           focus_steps=make_focus_steps())
        names = wb.sheetnames
        assert names.index("Analiza") < names.index("Kolejnosc fokusa")
        assert names.index("Kolejnosc fokusa") < names.index("Legenda")

    def test_multiple_analysis_sheets(self):
        pages = [make_page(sheet_name="Strona 1"), make_page(sheet_name="Strona 2")]
        wb = make_workbook(pages)
        assert "Strona 1" in wb.sheetnames
        assert "Strona 2" in wb.sheetnames

    def test_returns_bytes(self):
        gen = XLSXReportGenerator()
        result = gen.generate_bytes([make_page()])
        assert isinstance(result, bytes)
        assert result[:2] == b"PK"  # ZIP signature

    def test_generate_file_saves_to_path(self, tmp_path):
        path = str(tmp_path / "test_audyt.xlsx")
        gen = XLSXReportGenerator()
        returned = gen.generate_file([make_page()], path=path)
        assert returned == path
        with open(path, "rb") as f:
            assert f.read(2) == b"PK"


# ---------------------------------------------------------------------------
# Testy: arkusz analizy — naglowki i struktura
# ---------------------------------------------------------------------------

class TestAnalysisSheetStructure:
    def test_row1_is_title(self):
        wb = make_workbook([make_page(title="Audyt strony")])
        ws = wb.active or wb.worksheets[0]
        assert ws["A1"].value == "Audyt strony"

    def test_row3_headers_correct(self):
        wb = make_workbook([make_page()])
        ws = wb.worksheets[0]
        for col_idx, expected in enumerate(COLUMN_HEADERS, 1):
            assert ws.cell(row=3, column=col_idx).value == expected

    def test_exactly_11_columns(self):
        wb = make_workbook([make_page()])
        ws = wb.worksheets[0]
        used_cols = {c.column for row in ws.iter_rows(min_row=3, max_row=3) for c in row if c.value}
        assert len(used_cols) == 11

    def test_header_row_background(self):
        wb = make_workbook([make_page()])
        ws = wb.worksheets[0]
        # Kolor tla naglowka A3
        hex_col = cell_fill_hex(ws["A3"])
        assert hex_col.upper() == C_HEADER_BG.upper()

    def test_freeze_panes_at_A4(self):
        wb = make_workbook([make_page()])
        ws = wb.worksheets[0]
        assert ws.freeze_panes == "A4"

    def test_auto_filter_on_row3(self):
        wb = make_workbook([make_page()])
        ws = wb.worksheets[0]
        assert ws.auto_filter.ref is not None
        assert ws.auto_filter.ref.startswith("A3")

    def test_data_starts_at_row4(self):
        page = make_page(issues=[
            Issue(id="i1", wcag_criterion="1.1.1 Tresc", severity="critical",
                  description="Opis", recommendation="Naprawa",
                  affected_elements=[])
        ])
        wb = make_workbook([page])
        ws = wb.worksheets[0]
        # Wiersz 4 = separator lub dane — powinien miec wartosc
        assert ws.cell(row=4, column=1).value is not None


# ---------------------------------------------------------------------------
# Testy: kolory statusow i poziomow
# ---------------------------------------------------------------------------

class TestColors:
    @pytest.mark.parametrize("status,hex_color", list(STATUS_COLORS.items()))
    def test_status_color_in_col_B(self, status, hex_color):
        issues = [Issue(id="i1", wcag_criterion="1.1.1 Tresc",
                        severity="critical" if status == "KRYTYCZNE" else "major",
                        description="D", recommendation="R", affected_elements=[])]
        page = IssueMapper().map_page(issues=issues, sheet_name="T", prefix="T")
        wb = make_workbook([page])
        ws = wb.worksheets[0]
        # Znajdz wiersz z danych (pomijamy separatory)
        for row in ws.iter_rows(min_row=4):
            b_cell = row[1]  # kolumna B (index 1)
            if b_cell.value in ("KRYTYCZNE", "NOWE", "WSPOLNE", "OK"):
                if b_cell.value == status:
                    fill_hex = cell_fill_hex(b_cell)
                    assert fill_hex.upper() == hex_color.upper()
                    break

    def test_level_A_color_in_col_F(self):
        issues = [Issue(id="i1", wcag_criterion="1.1.1 Tresc",
                        severity="critical", description="D",
                        recommendation="R", affected_elements=[])]
        page = IssueMapper().map_page(issues=issues, sheet_name="T", prefix="T")
        wb = make_workbook([page])
        ws = wb.worksheets[0]
        for row in ws.iter_rows(min_row=4):
            f_cell = row[5]  # kolumna F (index 5)
            if f_cell.value == "A":
                assert cell_fill_hex(f_cell).upper() == LEVEL_COLORS["A"].upper()
                break

    def test_level_AA_color_in_col_F(self):
        issues = [Issue(id="i1", wcag_criterion="1.4.3 Kontrast",
                        severity="major", description="D",
                        recommendation="R", affected_elements=[])]
        page = IssueMapper().map_page(issues=issues, sheet_name="T", prefix="T")
        wb = make_workbook([page])
        ws = wb.worksheets[0]
        for row in ws.iter_rows(min_row=4):
            f_cell = row[5]
            if f_cell.value == "AA":
                assert cell_fill_hex(f_cell).upper() == LEVEL_COLORS["AA"].upper()
                break


# ---------------------------------------------------------------------------
# Testy: czcionki kodu
# ---------------------------------------------------------------------------

class TestCodeFonts:
    def test_column_J_uses_consolas(self):
        issues = [Issue(id="i1", wcag_criterion="1.1.1 Tresc", severity="critical",
                        description="D", recommendation="R",
                        affected_elements=[AffectedElement(tag="img",
                            html_snippet="<img src='x'>", html_after="")])]
        page = IssueMapper().map_page(issues=issues, sheet_name="T", prefix="T")
        wb = make_workbook([page])
        ws = wb.worksheets[0]
        for row in ws.iter_rows(min_row=4):
            j_cell = row[9]  # kolumna J (index 9)
            if j_cell.value and j_cell.value.startswith("<"):
                assert j_cell.font.name == "Consolas"
                break

    def test_column_K_uses_consolas(self):
        issues = [Issue(id="i1", wcag_criterion="1.1.1 Tresc", severity="critical",
                        description="D", recommendation="R",
                        affected_elements=[AffectedElement(tag="img",
                            html_snippet="<img>", html_after='<img alt="ok">')])]
        page = IssueMapper().map_page(issues=issues, sheet_name="T", prefix="T")
        wb = make_workbook([page])
        ws = wb.worksheets[0]
        for row in ws.iter_rows(min_row=4):
            k_cell = row[10]  # kolumna K (index 10)
            if k_cell.value and k_cell.value.startswith("<"):
                assert k_cell.font.name == "Consolas"
                break

    def test_column_A_uses_arial(self):
        page = make_page()
        wb = make_workbook([page])
        ws = wb.worksheets[0]
        for row in ws.iter_rows(min_row=4):
            a_cell = row[0]  # kolumna A
            if a_cell.value and "-R" in str(a_cell.value):
                assert a_cell.font.name == "Arial"
                break


# ---------------------------------------------------------------------------
# Testy: separatory sekcji
# ---------------------------------------------------------------------------

class TestSeparatorRows:
    def test_separator_text_not_empty(self):
        page = make_page()
        wb = make_workbook([page])
        ws = wb.worksheets[0]
        sep_values = []
        for row in ws.iter_rows(min_row=4):
            if row[0].value and isinstance(row[0].value, str) and "POZIOM" in row[0].value.upper():
                sep_values.append(row[0].value)
        assert len(sep_values) > 0

    def test_separator_fill_color(self):
        page = make_page()
        wb = make_workbook([page])
        ws = wb.worksheets[0]
        from app.services.xlsx_report_generator import C_SEPARATOR_BG
        for row in ws.iter_rows(min_row=4):
            val = row[0].value
            if val and isinstance(val, str) and "POZIOM" in val.upper():
                assert cell_fill_hex(row[0]).upper() == C_SEPARATOR_BG.upper()
                break


# ---------------------------------------------------------------------------
# Testy: arkusz Kolejnosc fokusa
# ---------------------------------------------------------------------------

class TestFocusOrderSheet:
    def test_title_row(self):
        wb = make_workbook([make_page()], focus_steps=make_focus_steps())
        ws = wb["Kolejnosc fokusa"]
        assert ws["A1"].value is not None
        assert "fokusa" in ws["A1"].value.lower()

    def test_headers_row3(self):
        wb = make_workbook([make_page()], focus_steps=make_focus_steps())
        ws = wb["Kolejnosc fokusa"]
        for col_idx, expected in enumerate(FOCUS_COLUMN_HEADERS, 1):
            assert ws.cell(row=3, column=col_idx).value == expected

    def test_exactly_4_columns(self):
        wb = make_workbook([make_page()], focus_steps=make_focus_steps())
        ws = wb["Kolejnosc fokusa"]
        used = {c.column for row in ws.iter_rows(min_row=3, max_row=3) for c in row if c.value}
        assert len(used) == 4

    def test_pre_step_in_data(self):
        wb = make_workbook([make_page()], focus_steps=make_focus_steps())
        ws = wb["Kolejnosc fokusa"]
        pre_cells = [ws.cell(row=r, column=1).value
                     for r in range(4, ws.max_row + 1)]
        assert "Pre" in pre_cells

    def test_step_numbers_present(self):
        wb = make_workbook([make_page()], focus_steps=make_focus_steps())
        ws = wb["Kolejnosc fokusa"]
        values = [ws.cell(row=r, column=1).value
                  for r in range(4, ws.max_row + 1)]
        assert "1" in values
        assert "2" in values

    def test_empty_steps_shows_placeholder(self):
        wb = make_workbook([make_page()], focus_steps=[])
        ws = wb["Kolejnosc fokusa"]
        all_values = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
        assert any(v and "Brak danych" in str(v) for v in all_values)

    def test_freeze_panes(self):
        wb = make_workbook([make_page()], focus_steps=make_focus_steps())
        ws = wb["Kolejnosc fokusa"]
        assert ws.freeze_panes == "A4"

    def test_row_count_matches_steps(self):
        steps = make_focus_steps()
        wb = make_workbook([make_page()], focus_steps=steps)
        ws = wb["Kolejnosc fokusa"]
        # 3 naglowki + 1 separator + len(steps) wierszy danych
        assert ws.max_row == 3 + 1 + len(steps)


# ---------------------------------------------------------------------------
# Testy: arkusz Legenda
# ---------------------------------------------------------------------------

class TestLegendSheet:
    def test_legend_title(self):
        wb = make_workbook([make_page()])
        ws = wb["Legenda"]
        assert ws["A1"].value is not None
        assert "Legenda" in ws["A1"].value

    def test_legend_has_status_entries(self):
        wb = make_workbook([make_page()])
        ws = wb["Legenda"]
        all_values = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
        for status in ("KRYTYCZNE", "NOWE", "WSPOLNE", "OK"):
            assert any(status in str(v) for v in all_values if v)

    def test_legend_has_level_entries(self):
        wb = make_workbook([make_page()])
        ws = wb["Legenda"]
        all_values = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
        assert any("A" == str(v).strip() for v in all_values if v)
        assert any("AA" == str(v).strip() for v in all_values if v)

    def test_legend_has_column_descriptions(self):
        wb = make_workbook([make_page()])
        ws = wb["Legenda"]
        all_values = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
        assert any("ID" in str(v) for v in all_values if v)
        assert any("Fragment kodu" in str(v) for v in all_values if v)

    def test_legend_min_rows(self):
        wb = make_workbook([make_page()])
        ws = wb["Legenda"]
        # 2 naglowki + 3 sekcje + 4 statusy + 3 poziomy + 11 kolumn = min 23
        assert ws.max_row >= 23

    def test_legend_3_columns(self):
        wb = make_workbook([make_page()])
        ws = wb["Legenda"]
        used = {c.column for row in ws.iter_rows(min_row=2, max_row=2) for c in row if c.value}
        assert len(used) == 3
